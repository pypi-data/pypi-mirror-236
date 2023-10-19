import openpyxl as xl
import pandas as pd
from deprecated import deprecated
import os
from io import BytesIO
from tempfile import NamedTemporaryFile
import re


def workbook_to_bytes(wb: xl.workbook.workbook.Workbook, lock_worksheets=False):
    if lock_worksheets:
        for x in wb.worksheets:
            x.protection.enable() # always do this before exporting because they automatically become unprotected somehow

    # If on windows, add delete=False to NamedTemporaryFile()
    ntf_delete = False if os.name == 'nt' else True
    with NamedTemporaryFile(delete=ntf_delete) as tmp:
        wb.save(tmp.name)
        excel_bytes = BytesIO(tmp.read())
        excel_bytes.seek(0)
        excel_bytes = excel_bytes.read()
    return excel_bytes

def df_to_worksheet(df: pd.DataFrame, cell: xl.cell.cell.Cell):
    ws = cell.parent

    for r in range(df.shape[0]):
        for c in range(df.shape[1]):
            cell.offset(r+1, c).value = df.iloc[r, c]

@deprecated(version='0.3.0', reason='renamed to df_to_worksheet' )
def df_to_sh(df: pd.DataFrame, cell: xl.cell.cell.Cell):
    df_to_worksheet(df, cell)

def df_to_excel_table_resize(df, wb, sheet_name, table_name):
    data_db = wb[sheet_name]
    assert table_name in data_db.tables, f"Table {table_name} does not exist in the worksheet {sheet_name}"
    ref_start, ref_end = data_db.tables[table_name].ref.split(':')

    start_col, start_row = split_string(ref_start)
    end_col, end_row = split_string(ref_end)
    table_length = (excel_letter_to_numeric(end_col))-(excel_letter_to_numeric(start_col)-1)

    assert table_length == len(df.columns), f"Dataframe width does not match table: df has {len(df.columns)} columns and {table_name} has {table_length} columns."

    # Blank out range (if data already exists)
    # TODO: Add check for keeping headers?
    for row in data_db[f"{start_col}{start_row+1}:{end_col}{end_row}"]:
        for cell in row:
            cell.value = None

    df_to_worksheet(df, data_db[f"{start_col}{start_row}"])
    data_db.tables[table_name].ref = f"{start_col}{start_row}:{end_col}{start_row+len(df)}"

def get_named_range_cell(wb, rng_name):
    '''return cell object from range name'''
    r = list(wb.defined_names[rng_name].destinations)[0]
    return wb[r[0]][r[1]]


def letter_to_number(letter):
    return ord(letter.upper()) - ord('A') + 1

def excel_letter_to_numeric(letters):
    result = 0
    for letter in letters:
        result = result * 26 + letter_to_number(letter)
    return result

def number_to_letter(number):
    return chr(number + ord('A') - 1)

def numeric_to_excel_letter(number):
    result = ''
    while number > 0:
        number, remainder = divmod(number - 1, 26)
        result = number_to_letter(remainder + 1) + result
    return result


def split_string(string):
    col, row = re.findall(r'\D+|\d+', string)
    return [col, int(row)]